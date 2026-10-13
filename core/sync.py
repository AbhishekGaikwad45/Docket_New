"""Employee master import, ported from Docket's sync_service.

Same rules as Docket:
- column names are matched loosely ("EMPLOYEE  NAME", "EMPLOYEE_NAME", "Name" all work)
- emails lower-cased; a blank email from the source never wipes one saved in the portal
- department "IT" is renamed "Information Technology"
- staff records win over associates when the same ID appears in both
Addition: a reporting manager column is picked up when the source has one.
Instead of purging employees missing from the feed, they can be marked Inactive (keeps history intact).
"""
import csv
import io
import re
from datetime import datetime

from core.db import db
from core.models import Employee

COLUMNS = {
    "employee_id": ("EMPLOYEE ID", "EMPID", "EMP ID", "EMP CODE", "EMPLOYEE CODE"),
    "name": ("EMPLOYEE NAME", "NAME"),
    "designation": ("DESIGNATION",),
    "department": ("DEPARTMENT", "DEPT"),
    "email": ("EMAIL ID", "EMAIL", "OFFICIAL EMAIL"),
    "status": ("EMPLOYEE STATUS", "STATUS"),
    "phone": ("CONTACT NO", "MOBILE", "MOBILE NO", "PHONE"),
    "gender": ("GENDER",),
    "category": ("CATEGORY",),
    "manager_id": ("REPORTING MANAGER ID", "REPORTING MANAGER CODE", "MANAGER ID", "RM ID",
                   "REPORTING MANAGER", "REPORTING TO", "REPORTING OFFICER"),
}
DEPARTMENT_RENAMES = {"IT": "Information Technology"}
SOURCE_PRIORITY = {"manual": -1, "staff": 0, "associates": 1, "upload": 2}


def _key(name):
    return re.sub(r"[\s_.]+", "", str(name or "").upper())


def _clean(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():  # Excel turns 4050163 into 4050163.0
        value = int(value)
    text = str(value).strip()
    return text or None


def normalize(header, rows, source):
    """Map raw rows from any source to employee dicts."""
    keys = [_key(h) for h in header]
    idx = {}
    for field, names in COLUMNS.items():
        for name in names:
            if _key(name) in keys:
                idx[field] = keys.index(_key(name))
                break
    if "employee_id" not in idx:
        raise ValueError("No employee ID column found. Expected one of: " + ", ".join(COLUMNS["employee_id"]))

    records = {}
    for row in rows:
        get = lambda f: _clean(row[idx[f]]) if f in idx and idx[f] < len(row) else None  # noqa: E731
        emp_id = get("employee_id")
        if not emp_id:
            continue
        dept = get("department")
        manager = get("manager_id")
        if manager:  # "Ravi Kumar (4050163)" -> "4050163"
            found = re.search(r"\(([^()]+)\)\s*$", manager)
            manager = found.group(1).strip() if found else manager
        records[emp_id] = {
            "employee_id": emp_id,
            "name": get("name") or f"Employee {emp_id}",
            "designation": get("designation"),
            "department": DEPARTMENT_RENAMES.get((dept or "").upper(), dept),
            "email": (get("email") or "").lower() or None,
            "status": get("status") or "Active",
            "phone": get("phone"),
            "gender": get("gender"),
            "category": get("category"),
            "manager_id": manager,
            "source": source,
        }
    return list(records.values())


def upsert(records, deactivate_missing=False):
    """Insert/update employees. Returns a stats dict. Caller commits."""
    now = datetime.now()
    stats = {
        "received": len(records),
        "inserted": 0,
        "updated": 0,
        "kept_staff": 0,
        "kept_manual": 0,
        "managers_linked": 0,
        "deactivated": 0,
    }
    existing = {e.employee_id: e for e in Employee.query.all()}

    for rec in records:
        emp = existing.get(rec["employee_id"])
        if emp is None:
            emp = Employee(employee_id=rec["employee_id"], created_at=now)
            db.session.add(emp)
            existing[emp.employee_id] = emp
            stats["inserted"] += 1
        elif emp.is_superadmin or emp.source == "manual" or emp.employee_id == "admin":
            stats["kept_manual"] += 1
            continue  # Admin or manually edited in portal stays as is!
        elif SOURCE_PRIORITY.get(emp.source, 9) < SOURCE_PRIORITY.get(rec["source"], 9):
            stats["kept_staff"] += 1  # e.g. associates feed must not overwrite a staff record
            continue
        else:
            stats["updated"] += 1

        for field in ("name", "status", "phone", "gender", "category", "source"):
            val = rec[field]
            if val is not None and str(val).strip() != "":
                setattr(emp, field, val)
        if rec["designation"]:
            emp.designation = rec["designation"]
        if rec["department"]:
            emp.department = rec["department"]
        emp.email = rec["email"] or emp.email
        emp.last_synced_at = now

    # Second pass: managers may appear later in the feed than their reports.
    for rec in records:
        emp = existing[rec["employee_id"]]
        if (
            rec["manager_id"]
            and not emp.manager_locked
            and not emp.is_superadmin
            and emp.source != "manual"
            and rec["manager_id"] in existing
            and rec["manager_id"] != emp.employee_id
            and emp.source == rec["source"]
        ):
            emp.reporting_manager_id = rec["manager_id"]
            stats["managers_linked"] += 1

    if deactivate_missing and records:
        seen = {r["employee_id"] for r in records}
        sources = {r["source"] for r in records}
        for emp in existing.values():
            if emp.source in sources and emp.employee_id not in seen and emp.is_active and not emp.is_superadmin:
                emp.status = "Inactive"
                stats["deactivated"] += 1
    db.session.flush()
    return stats


def read_file(filename, stream):
    """(header, rows) from an uploaded .xlsx or .csv."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        sheet = load_workbook(stream, read_only=True, data_only=True).active
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    elif filename.lower().endswith(".csv"):
        text = stream.read().decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(text)))
    else:
        raise ValueError("Upload an .xlsx or .csv file.")
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        raise ValueError("The file is empty.")
    return rows[0], rows[1:]


def fetch_view(odbc, view, active_only=False):
    """(header, rows) from a SQL Server view (Mantra)."""
    import pyodbc  # optional dependency, only needed for live sync
    if not re.fullmatch(r"\w+", view):
        raise ValueError(f"Bad view name {view!r}")
    query = f"SELECT * FROM dbo.{view}" + (" WHERE [Employee Status] = 'Active'" if active_only else "")
    with pyodbc.connect(odbc, timeout=15) as conn:
        cur = conn.cursor()
        cur.execute(query)
        return [c[0].strip() for c in cur.description], [list(r) for r in cur.fetchall()]
